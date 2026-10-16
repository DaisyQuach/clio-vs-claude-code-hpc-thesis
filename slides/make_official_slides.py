#!/usr/bin/env python3
"""Build the OFFICIAL thesis deck: CLIO (clio-agent 0.9.5b5.post2 + factorio-flat 0.5.0) vs the Claude Code baseline,
n = 5 counted trials per experiment (E1, E2, E3). Same structure as the pilot deck (make_hpc2_slides.py).

Rebuild:
    1. python3 official_thesis_clio/official_aggregate.py     (both arms read-only from Thesis/)
    2. python3 official_thesis_clio/official_compare.py       (data/, figures/H1–H6)
    3. slides/.venv/bin/python slides/make_official_slides.py -> slides/CLIO_vs_Claude_Code_official_thesis.pptx
"""
import collections
import csv
import datetime as dt
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

B = "/uufs/chpc.utah.edu/common/home/chet/Daisy/clioSessions/CLIO_E1_API_pilots"
HV = f"{B}/official_thesis_clio"
FIG, DATA = f"{HV}/figures", f"{HV}/data"
OUT = f"{B}/slides/CLIO_vs_Claude_Code_official_thesis.pptx"
INK, INK2, ACC, SURF = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x62, 0x50, 0xD6), RGBColor(0xFC, 0xFC, 0xFB)
RED = RGBColor(0xE3, 0x49, 0x48)
LABEL = "COUNTED THESIS TRIALS · CLIO (factorio-flat 0.5.0) vs Claude Code baseline · n = 5 per system per experiment"
EXPS = ("E1", "E2", "E3")

NARRATIVE = {
    "question": [
        "Does CLIO (clio-agent with the factorio-flat 0.5.0 blueprint) reach the same scientific outcome as the Claude Code baseline, and at what cost and time?",
        "Counted thesis trials: five per system per experiment. E1 uninterrupted, E2 interrupted recovery, E3 disparate storage.",
        "Measured: scored outcome (thesis scorer), cost, wall/model/tool/Slurm time, API calls, tokens, Morphorm jobs and restarts.",
    ],
    "design": [
        "Baseline: the 5 counted Claude Code trials per experiment (Claude Code 2.1.295, claude-fable-5-1, CLI-default effort, --dangerously-skip-permissions).",
        "CLIO: 5 counted trials per experiment, each a fresh deployment in its own agent-host job with the baseline's Slurm contract: clio-agent 0.9.5b5.post2, "
        "provider claude_code driving the same Claude Code 2.1.295 binary and model, factorio-flat 0.5.0, approval bypass, CLI-default effort, "
        "runtime tuning 3600/3600/5400 s, private clio-core daemon, no user-configured MCP servers, stream audit on. The blueprint's own web MCP (one tool, web_fetch) was mounted; like the baseline's built-in WebSearch/WebFetch it was forbidden by the task rules and never called.",
        "Identical for both: frozen packets (FROZEN-EX8-v1 / FROZEN-EX16-v1 / E3 control packet), byte-identical prompts, one user message per session (no nudges), "
        "6600 s session limit, same scorer, acceptance specs and tolerances, same packet verification.",
        "Trials ran one at a time (as the baseline did). Statistics: ratio of means with bootstrap 95 % CI; exact permutation test (252 splits) and exact "
        "Mann-Whitney U. 'Different' only when p < 0.05 and the CI excludes 1. With n = 5, 'no detectable difference' is not proof of equality.",
    ],
    "caveats": [
        "The CLIO trials ran on 2026-10-13, three to four days after the baseline (2026-10-09/10), on the same partition; node and contention differ between runs.",
        "Slurm time depends on node and load (the same Morphorm job takes about 10 min on an idle notch351 and 17–19 min on notch272–274). "
        "Compare 'time outside Slurm' for agent speed.",
        "factorio-flat 0.5.0 contains generic Morphorm/Slurm guidance and helper scripts that the baseline agent did not have; two items were informed by "
        "non-counted pilots (filter keyword form, restart-state counting). The blueprint is part of the CLIO system under test.",
        "factorio-flat 0.5.0 here is the frozen pilot-validated copy (provenance and checksum: Thesis/automation/clio/blueprints/factorio-flat-0.5.0.PROVENANCE.md).",
        "CLIO's context, assimilation, transfer and exploration engines are not exercised by E1–E3 (see 'Which parts of CLIO did E1–E3 exercise?'); results reflect the agent harness, blueprint and runtime.",
        "E3 boundary audit (researcher decision 2026-10-13): reads of the CLIO agent's own runtime (its blueprint skills and its own session's saved tool output, 8–10 per trial) "
        "are re-labelled OWN_RUNTIME; every other finding still fails. The shared audit's own output (FAIL on those reads only) is kept in each trial. "
        "Counted CLIO E3 trials are 03–07 (06 and 07 replace the aborted 01 and the excluded 02).",
    ],
}


def rows(path):
    return [r for r in csv.DictReader(l for l in open(path) if not l.startswith("#"))] if os.path.exists(path) else []


PR = rows(f"{DATA}/per_run.csv")
ST = {(s["experiment"], s["metric"]): s for s in rows(f"{DATA}/stats.csv")}
MJ = rows(f"{DATA}/morphorm_jobs.csv")
f = lambda x: float(x) if x not in ("", "nan", None) else float("nan")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def base_slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = SURF
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8)).text_frame; tb.word_wrap = True
    tb.text = title; p = tb.paragraphs[0]; p.font.size = Pt(26); p.font.bold = True; p.font.color.rgb = INK
    if subtitle:
        q = tb.add_paragraph(); q.text = subtitle; q.font.size = Pt(13); q.font.color.rgb = INK2
    ft = s.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.3)).text_frame
    ft.text = LABEL; ft.paragraphs[0].font.size = Pt(9); ft.paragraphs[0].font.color.rgb = INK2
    return s


def bullets(s, items, top=1.5, left=0.6, width=12.1, size=17):
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(6.95 - top)).text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sub = t.startswith("    ")
        p.text = t.strip() if sub else "• " + t
        p.level = 1 if sub else 0
        p.font.size = Pt(size - 2 if sub else size); p.font.color.rgb = INK; p.space_after = Pt(6)


def picture(s, name, top=1.35, max_h=5.6, max_w=12.3):
    path = name if name.endswith(".png") else f"{FIG}/{name}.png"
    if not os.path.exists(path):
        bullets(s, [f"(figure {os.path.basename(path)} not generated yet)"]); return
    pic = s.shapes.add_picture(path, Inches(0.5), Inches(top))
    r = min(Inches(max_w) / pic.width, Inches(max_h) / pic.height)
    pic.width, pic.height = int(pic.width * r), int(pic.height * r)
    pic.left = int((prs.slide_width - pic.width) / 2)


def table(s, header, data, top=1.4, col_w=None, size=11, red_cells=()):
    nr, nc = len(data) + 1, len(header)
    t = s.shapes.add_table(nr, nc, Inches(0.5), Inches(top), Inches(12.3), Inches(min(5.5, 0.3 * nr))).table
    for j, v in enumerate(header):
        c = t.cell(0, j); c.text = v; c.text_frame.paragraphs[0].font.size = Pt(size); c.text_frame.paragraphs[0].font.bold = True
    for i, r in enumerate(data, 1):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v); p = c.text_frame.paragraphs[0]; p.font.size = Pt(size)
            if (i, j) in red_cells:
                p.font.color.rgb = RED; p.font.bold = True
    if col_w:
        for j, w in enumerate(col_w):
            t.columns[j].width = Inches(w)


def ms(e, sysn, m, fmt="{:.1f}"):
    v = [f(r[m]) for r in PR if r["experiment"] == e and r["system"] == sysn]
    if not v:
        return "—"
    mean = sum(v) / len(v); sd = (sum((x - mean) ** 2 for x in v) / (len(v) - 1)) ** 0.5 if len(v) > 1 else float("nan")
    return (fmt.format(mean) + (" ± " + fmt.format(sd) if sd == sd else ""))


def n(e, sysn):
    return sum(1 for r in PR if r["experiment"] == e and r["system"] == sysn)


def outcomes(e, sysn):
    c = collections.Counter(r["outcome"] for r in PR if r["experiment"] == e and r["system"] == sysn)
    return ", ".join(f"{k} {v}/{sum(c.values())}" for k, v in c.items()) or "—"


def verdict_line(e, m, lab, unit=""):
    s = ST.get((e, m))
    if not s:
        return f"{e} {lab}: no data"
    r, lo, hi, p = f(s["ratio"]), f(s["ratio_ci_lo"]), f(s["ratio_ci_hi"]), f(s["perm_p"])
    v = s["verdict"]
    word = {"lower": "LOWER for CLIO", "higher": "HIGHER for CLIO", "no detectable difference": "no detectable difference",
            "too few runs (n < 3)": "too few runs for a verdict"}[v]
    ci = f", 95 % CI {lo:.2f}–{hi:.2f}" if lo == lo else ""
    return f"{e} {lab}: {word} (×{r:.2f}{ci}; p = {p:.3f})" if p == p else f"{e} {lab}: ×{r:.2f} (n too small for a test)"


# ------------------------------------------------------------------ slides
NS = " / ".join(str(n(e, "clio")) for e in EXPS)
s = base_slide(f"CLIO vs Claude Code baseline: counted thesis trials (CLIO n = {NS} for E1/E2/E3; baseline n = 5)",
               f"E1 / E2 / E3 Morphorm packets · updated {dt.date.today().isoformat()}")
bullets(s, NARRATIVE["question"], top=2.0, size=21)

s = base_slide("What we compared", "Same frozen packets, prompts, model and agent-host contract; only the system around the model differs")
bullets(s, NARRATIVE["design"], size=15)

ATT = rows(f"{HV}/data/trial_attempts.csv")
s = base_slide("Trial attempts", "Same rule as the baseline: non-agent failures are kept, marked invalid and replaced with a new numbered attempt")
table(s, ["Exp", "Attempt", "Status", "Note"], [(a["experiment"], a["attempt"], a["status"], a["note"]) for a in ATT] or [("—", "—", "—", "no invalid attempts")],
      col_w=[0.8, 1.5, 2.0, 8.0], size=11)
s = base_slide("Headline: outcome, cost and time (mean ± SD)", "Wall/model/outside-Slurm minutes summed over sessions (E2 = A + B)")
data = []
for e in EXPS:
    for sysn, lab in (("baseline", "Claude Code baseline"), ("clio", "CLIO 0.5.0")):
        data.append((e, lab, n(e, sysn), outcomes(e, sysn), "$" + ms(e, sysn, "cost", "{:.2f}"), ms(e, sysn, "wall_min"), ms(e, sysn, "model_min"),
                     ms(e, sysn, "non_slurm_min"), ms(e, sysn, "api_calls", "{:.0f}"), ms(e, sysn, "cancelled_jobs", "{:.1f}")))
table(s, ["Exp", "System", "n", "Outcome", "Cost", "Wall min", "Model min", "Outside Slurm min", "API calls", "Cancelled jobs"], data,
      col_w=[0.5, 2.0, 0.4, 1.9, 1.4, 1.3, 1.3, 1.5, 1.1, 0.9], size=10)

s = base_slide("Same or different? Statistical comparison", "Ratio = CLIO 0.5.0 mean ÷ baseline mean; red = significant (p < 0.05 and CI excludes 1)")
MET = [("cost", "Cost"), ("wall_min", "Wall time"), ("model_min", "Model time"), ("non_slurm_min", "Outside Slurm"), ("slurm_min", "Slurm wait"),
       ("api_calls", "API calls"), ("output", "Output tokens"), ("cache_write", "Cache writes")]
data, red = [], []
for i, (m, lab) in enumerate(MET, 1):
    row = [lab]
    for j, e in enumerate(EXPS):
        st_ = ST.get((e, m))
        if not st_:
            row.append("—"); continue
        r, lo, hi, p = f(st_["ratio"]), f(st_["ratio_ci_lo"]), f(st_["ratio_ci_hi"]), f(st_["perm_p"])
        row.append(f"×{r:.2f} [{lo:.2f}–{hi:.2f}] p={p:.3f}" if p == p else f"×{r:.2f} (n<2)")
        if st_["verdict"] in ("lower", "higher"):
            red.append((i, j + 1))
    data.append(row)
table(s, ["Metric", "E1", "E2", "E3"], data, col_w=[2.1, 3.4, 3.4, 3.4], size=11, red_cells=red)
bullets(s, ["Tests: exact permutation test on the difference in means (also exact Mann-Whitney U, in data/stats.csv); bootstrap 95 % CI of the ratio. "
            "With n = 5, a non-significant result means 'no detectable difference', not 'equal'."], top=4.6, size=12)

s = base_slide("Verdict per experiment")
items = []
for e in EXPS:
    items.append(f"{e}: outcomes CLIO {outcomes(e, 'clio')} vs baseline {outcomes(e, 'baseline')}")
    for m, lab in (("cost", "cost"), ("wall_min", "wall time"), ("model_min", "model time"), ("non_slurm_min", "time outside Slurm"), ("api_calls", "API calls")):
        items.append("    " + verdict_line(e, m, lab))
bullets(s, items, size=14)

s = base_slide("Where the money goes", "Mean cost per run by component")
picture(s, "H1_cost_decomposition")
s = base_slide("Per-call view: context sent and cache writes", "Median-cost run of each system and experiment")
picture(s, "H2_per_call_context_and_cache_writes", top=1.2, max_h=5.65)
s = base_slide("Time: model minutes, tools and Slurm waits", "Slurm/Morphorm waits depend on node and contention; model and tool time are the agent's part")
picture(s, "H3_time_breakdown", top=1.3, max_h=5.6)
s = base_slide("How much work each run did", "Each dot is one run; bar = mean")
picture(s, "H4_work_volume", top=1.5, max_h=5.2)
s = base_slide("Per-run distributions: cost and time", "n = 5 vs n = 5; ratio of means and exact permutation p above each pair")
picture(s, "H5_per_run_distributions", top=1.3, max_h=5.65)
s = base_slide("Scorecard: CLIO 0.5.0 relative to the baseline", "Ratio of means with bootstrap 95 % CI; 1.0 = same")
picture(s, "H6_scorecard_vs_baseline", top=1.6, max_h=4.8)

s = base_slide("Morphorm jobs: nodes, run time and restarts", "Every Morphorm/Cubit job submitted by the agent (sacct)")
data = []
for e in EXPS:
    for sysn, lab in (("baseline", "Claude Code"), ("clio", "CLIO 0.5.0")):
        js = [m for m in MJ if m["experiment"] == e and m["system"] == sysn]
        if not js:
            continue
        done = [m for m in js if m["state"] == "COMPLETED"]
        nodes = collections.Counter(m["node"] for m in done)
        canc = sum(1 for m in js if m["state"].startswith("CANCELLED"))
        runs_with_cancel = len({m["run"] for m in js if m["state"].startswith("CANCELLED")})
        mean_el = sum(float(m["elapsed_min"]) for m in done) / max(1, len(done))
        data.append((e, lab, len({m["run"] for m in js}), len(js), canc, f"{runs_with_cancel}", f"{mean_el:.1f}",
                     ", ".join(f"{k}×{v}" for k, v in nodes.most_common())))
table(s, ["Exp", "System", "Runs", "Jobs", "Cancelled", "Runs with a cancel", "Mean completed job (min)", "Nodes (completed jobs)"], data,
      col_w=[0.5, 1.6, 0.7, 0.7, 1.0, 1.5, 1.9, 4.4], size=10.5)
bullets(s, ["E2 cancellations are the restart-index mistake (Stage 2 started from the wrong saved state, then cancelled and resubmitted).",
            "Job run time differs by node and load, which moves wall time independently of the agent."], top=5.6, size=13)

# ------------------------------------------------------------------ observability (same method as the pilot study)
OBS = f"{HV}/observability"
s = base_slide("Observability: which questions each system's own records answer",
               "Same 14 diagnostic questions and code-scored rules as the pilot study, now on the counted trials (n = 5 per system per experiment)")
picture(s, f"{OBS}/figures/obs1_diagnostic_coverage.png", top=1.25, max_h=5.75)

s = base_slide("Observability in the counted trials: numbers", "Only records each system writes about itself; files written by our runners are excluded")
table(s, ["", "E1", "E2", "E3", "Pilots (for comparison)"], [
    ("Claude Code baseline: points / 14 (fully recorded)", "8.5 (6)", "8.5 (6)", "8.5 (6)", "8.5 (6)"),
    ("CLIO with stream audit", "12 (11)", "12 (11)", "12 (11)", "11.5–12 (10)"),
    ("CLIO without stream audit (opt-in)", "9 (8)", "9 (8)", "9 (8)", "8.5–9 (7)"),
    ("Distinct record types per run: baseline vs CLIO", "31 vs 103", "31 vs 103", "31–35 vs 103", "31 vs 73–80"),
    ("Structured records per model call", "18–29 vs 188–246", "21–24 vs 217–230", "15–18 vs 213–257", "15–29 vs 180–244"),
    ("Disk for own records per run (CLIO + full exports)", "2.6 vs 21 MB (323 MB)", "4.0 vs 37 MB (849 MB)", "2.5 vs 23 MB (347 MB)", "≈2–4 vs 70–150 MB"),
], top=1.35, col_w=[4.1, 2.0, 2.0, 2.0, 2.2], size=11)
bullets(s, [
    "Same picture as the pilots, in all three experiments: CLIO answers 12 of 14 questions vs 8.5 for Claude Code; without the opt-in audit the two are close (9 vs 8.5) but cover different questions.",
    "Change from the pilots: Q4 (liveness of a long tool) is now 'recorded' for CLIO because the counted runner kept CLIO's SSE event stream (tool.call.progress); the pilots did not keep it.",
    "CLIO still lacks Q10 (running rate-limit state): it consumes the Claude Code stream and does not keep it. The baseline lacks Q2, Q12, Q13 (cache-reset reasons, prefix hashes, runtime health).",
    "Cost of the breadth: about 8–9× the storage for the records themselves, about 120–210× with CLIO's full session exports.",
], top=4.3, size=13)

s = base_slide("Observability in the counted trials: what it showed", "Cases where CLIO's own records answered a question about a counted trial")
bullets(s, [
    "Cache stayed warm in every counted CLIO session (Q2): the stream audit typed all 535 model calls; 515 were incremental ('delta') and 20 were full sends, "
    "one first call per session (5 + 10 + 5 sessions). Zero 'session_evicted' re-sends, which confirms that the idle-reaper fix from the pilots held.",
    "E2-05 exclusion (Q5): CLIO recorded Session B's failure as a typed ClaudeCodeSignedOutError with 0 tokens, so the cause (expired login, not the agent) "
    "was clear from CLIO's own trace.",
    "E1-03 exclusion (Q6/Q13): CLIO's server log recorded the failed version probe of the pinned CLI ([Errno 14] Bad address) and its fallback to 2.1.296; "
    "the CLI-pin check caught it, and the log showed why.",
    "Which skills the agent used (Q8): every counted CLIO session recorded skill.loaded twice (the blueprint's packet-mode skills); the baseline has no such event.",
    "Effective configuration (Q6): model, approval mode, effort, blueprint id and version, and timeouts were read back from the running CLIO for every session "
    "and frozen in configuration_freeze.json before the prompt was sent.",
    "Limits: neither system attributes Slurm jobs to the session (Q9) or records which call wrote each output file (Q14); both are reconstructed from tool output.",
], size=14)

s = base_slide("Which parts of CLIO did E1–E3 exercise?", "From each session's recorded toolset, SSE tool calls and stream audit (15 counted trials)")
table(s, ["CLIO component", "Used in E1–E3?", "Evidence"], [
    ("clio-agent: GACT agent loop, provider claude_code, blueprint (factorio-flat 0.5.0)", "yes", "every session; this is what was compared"),
    ("Tools", "shell_bash, load_skill only", "45 tools offered; all 20 counted sessions called only shell_bash (25–55 per session) and load_skill (2)"),
    ("clio-core: session/context store (events, transcript segments)", "yes, passively", "store.put / store.write rows in the audit; private daemon per session"),
    ("Context engine (compaction, recall_context, memory search)", "no", "no session approached the context limit (compaction.auto_skipped); recall/memory tools never called"),
    ("Assimilation engine (CAE: ingest files into tagged blobs)", "no", "no tool exposes it to the agent; not configured in the deployment"),
    ("Transfer engine (CTE: tiered RAM/disk storage, FUSE)", "no", "Morphorm and E3 staging used plain files, Slurm and Globus as in the baseline"),
    ("Exploration engine (CEE: query/retrieve)", "no", "no tool exposes it to the agent"),
], top=1.35, col_w=[4.3, 2.2, 5.8], size=11)
bullets(s, [
    "So the results measure CLIO as an agent harness (loop, blueprint, runtime, telemetry), not its data engines.",
    "The engines were tested separately, without a model (CLIO_engine_pilots/): transfer works but is slower than plain local storage when durability is matched and loses data on a crash; "
    "assimilation stores raw bytes without parsing Exodus; exploration has no semantic search in this build (indexer chimod missing).",
], top=5.3, size=13)

s = base_slide("What the experiments show (qualitative)", "Same model and CLI in both arms; only the agent system around them differs")
table(s, ["Experiment", "What CLIO changed", "Why"], [
    ("All three", "Same accepted science, no human input; less writing, fewer and larger model calls, much richer self-records",
     "The harness and blueprint change how the model works, not what it can solve"),
    ("E1 uninterrupted", "Went straight to a correct deck and job; output 0.68×, cost 0.72×; agent time n.s.",
     "Blueprint Morphorm/Slurm guidance and bundled checks; wall time higher only from Slurm node speed"),
    ("E2 interrupted recovery", "Clearest advantage: no wrong-restart cancels (baseline 4/5 runs); model time 0.73×, cost 0.78×",
     "Encoded restart-state knowledge, not memory: Session B was a fresh deployment, as in the baseline"),
    ("E3 disparate storage", "Cheaper (0.83×), calls 0.55×; time and context ≈ baseline; staging behaviour the same",
     "Only the Morphorm half was helped; nothing in CLIO's blueprint or tools addressed finding and staging data"),
], top=1.4, col_w=[2.0, 5.3, 5.0], size=12)
bullets(s, ["Thesis point: CLIO's advantage scales with how much of the task its encoded domain knowledge covers."], top=5.5, size=15)

s = base_slide("Why E3 improved least", "Not because E3 relied more on clio-core: the core was used as little in E3 as in E1")
bullets(s, [
    "No blueprint help for the data half: every E3 session loaded the same two skills as E1 (morphorm, hpc-batch-jobs). factorio-flat 0.5.0 has no skill "
    "for storage discovery, relevance judgement, Globus staging or provenance.",
    "Data engines not in play: the agent called only shell_bash and load_skill; staging used Globus and plain files exactly as the baseline "
    "(14.4 vs 14.8 files staged, 11.6 vs 13.2 Globus transfers per run).",
    "Reading that cannot be skipped: cache writes (≈ half the cost) track the final context. E1: CLIO 161k vs 221k tokens → cache-write cost 0.77×. "
    "E3: 183k vs 187k → 1.04×. CLIO's first prompt (≈22.5k vs ≈20.7k tokens) is also written to cache fresh every session, "
    "while the baseline's first call reads most of its prompt from a cache warmed by earlier sessions.",
    "Fewer calls, same work: 27 vs 49 calls, but ≈6.2k vs ≈3.5k new tokens per call, so the saving shows in cache reads (0.54×) and cost, not time.",
    "Tool-layer overhead where calls are many and small: E3 tool time 3.2 vs 1.7 min (reading and record-keeping around staging).",
    "So E3 does not show CLIO is weak at data movement; it shows that part of CLIO was not tested.",
], size=14)

s = base_slide("Next: E3 with the clio-core engines", "Three levels, cheapest first; Levels 1–2 piloted (non-counted, Sonnet), Level 3 not run")
table(s, ["Level", "What changes", "What it tests", "Expected"], [
    ("1. Transparent transfer engine", "FUSE mount over a private disk-only CTE tier at the existing $E3_STAGING_DIR path; packet, prompt and helpers unchanged",
     "Does E3 run unchanged on the engine?", "Piloted: yes, accepted; no change in time"),
    ("2. Agent-visible engines", "MCP tools over the 2.3.1 engines (BM25 search of staged files, ingest, read, manifest) + a staging/provenance skill",
     "Does an engine-aware agent stage and find data more efficiently?", "Piloted: tools never called (0 in 2 runs); skill alone halved files staged"),
    ("3. Engine-dependent", "Parsed assimilation (Exodus/HDF5) + semantic search (indexer chimod)",
     "Can the agent query data instead of reading it?", "Blocked: missing in iowarp-core 2.2.1"),
], top=1.4, col_w=[2.3, 4.4, 3.0, 2.6], size=11.5)
bullets(s, ["Level 1 (2026-10-13): both arms accepted; same wall and Morphorm job time; staged files intact after clean stop + recover → compatible, no benefit at ≈2 MB.",
            "Level 2 (2026-10-14): all runs accepted; the agent made 0 engine-tool calls in 2 of 2 runs (kept grep/sed); the staging skill alone halved files staged "
            "and Globus transfers. Offering the engine is not using it: knowledge changed behaviour, tools did not.",
            "Constraints: crash loses files (2.2.1 pilot: 0/237 after kill -9) → clean unmount and stop; the agent's own store stays 2.2.1 (clio-agent pins ~=2.2.1)."],
        top=5.0, size=12.5)

s = base_slide("Caveats")
bullets(s, NARRATIVE["caveats"], size=15)

for sysn, lab in (("baseline", "Claude Code baseline"), ("clio", "CLIO 0.5.0")):
    s = base_slide(f"Per-run table: {lab}", "Minutes summed over sessions; Canc. = cancelled Morphorm jobs; Nudges = extra user messages (none allowed in counted trials)")
    data = [(r["experiment"], r["run"], r["outcome"], f"${f(r['cost']):.2f}", f"{f(r['wall_min']):.1f}", f"{f(r['model_min']):.1f}",
             f"{f(r['non_slurm_min']):.1f}", f"{f(r['slurm_min']):.1f}", r["api_calls"].split(".")[0], f"{f(r['output']) / 1e3:.0f}k",
             r["cancelled_jobs"].split(".")[0], r["nudges"].split(".")[0], r["nodes"]) for r in PR if r["system"] == sysn]
    table(s, ["Exp", "Run", "Outcome", "Cost", "Wall", "Model", "Out. Slurm", "Slurm", "Calls", "Output", "Canc.", "Nudges", "Morphorm node(s)"], data,
          col_w=[0.5, 0.7, 1.1, 0.9, 0.8, 0.8, 1.0, 0.8, 0.7, 0.8, 0.6, 0.7, 2.9], size=9.5)

s = base_slide("Appendix: sources and how to rebuild")
bullets(s, [
    "Trials: Thesis/runs/<experiment>/{baseline,clio}/trial_NN; scores: Thesis/researcher_only/scores/<run_id>/ (researcher overrides in RESEARCHER_DECISIONS.md)",
    "CLIO runner: Thesis/automation/clio/run_counted_thesis_clio.sh (batch log in researcher_only/batch_logs/clio_counted_*.log); config: clio_trial_config.env",
    "Blueprint: Thesis/automation/clio/blueprints/factorio-flat-0.5.0 (+ .PROVENANCE.md, .SHA256SUMS.txt)",
    "Comparison data: clioSessions/CLIO_E1_API_pilots/official_thesis_clio/data/{per_run,stats,morphorm_jobs}.csv; method: official_compare.py docstring",
    "Rebuild: python3 official_thesis_clio/official_aggregate.py → python3 official_thesis_clio/official_compare.py → slides/.venv/bin/python slides/make_official_slides.py",
    "Observability: python3 official_thesis_clio/official_observability.py (→ official_thesis_clio/observability/{data,figures}; method as observability_study/OBSERVABILITY.md)",
    "Prices: claude-fable-5-1 list ($10/M input, $50/M output, $0.25/M cache read, $20/M 1-h cache write)",
], size=14)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides))
